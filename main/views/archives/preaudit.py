"""任前联审登记表"""

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.conf import settings
from main.utils import _get_conn
import json
import logging
import os
import openpyxl
from urllib.parse import quote
from main.utils.decorators import archive_perm_required
from io import BytesIO
from django.http import HttpResponse
# @archive_perm_required

logger = logging.getLogger(__name__)


def _get_yhbh(request):
    try:
        return request.user.profile.yhbh or 0
    except:
        return 0


@login_required
@archive_perm_required
def person_preaudit_view(request):
    return render(request, "archives/person_preaudit.html")


@login_required
def preaudit_data_api(request):
    """获取任前联审数据"""
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        yhbh = _get_yhbh(request)
        cursor.execute("{CALL ZXSH_RQLSDJB(?, 'SELECT', ?)}", (int(rsid), yhbh))
        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        data = dict(zip(cols, row)) if row else {}
        cursor.close()

        # 处理意见自动填充
        if data.get("rqls28", "") and len(data.get("rqls28", "")) <= 5:
            data["rqls29"] = "此表信息已核实确认无误，不影响任用。"
        else:
            data["rqls29"] = (
                data.get("rqls29", "")
                or "此表信息已核实确认无误，干部档案存在的问题不影响任用。"
            )

        # 布尔字段
        bool_fields = [
            "rqls3",
            "rqls4",
            "rqls5",
            "rqls6",
            "rqls7",
            "rqls8",
            "rqls9",
            "rqls10",
            "rqls11",
            "rqls22",
            "rqls23",
        ]
        for f in bool_fields:
            data[f] = bool(data.get(f))

        return JsonResponse({"code": 0, "data": data})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
def preaudit_export_api(request):
    """导出任前联审登记表"""
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    template_path = os.path.join(
        settings.BASE_DIR, "static", "excel_templates", "任前联审登记表.xlsx"
    )
    if not os.path.exists(template_path):
        return JsonResponse({"code": 500, "msg": "模板不存在"})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        yhbh = _get_yhbh(request)
        cursor.execute("{CALL ZXSH_RQLSDJB(?, 'SELECT', ?)}", (int(rsid), yhbh))
        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        data = dict(zip(cols, row)) if row else {}
        cursor.close()

        if not data:
            return JsonResponse({"code": 404, "msg": "请先填写专审情况登记表"})

        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        merged = list(ws.merged_cells.ranges)
        for mr in merged:
            ws.unmerge_cells(str(mr))

        def yn(v, default=1):
            if default == 0 and not v:
                return "是 ☐  否 ☐"
            return "是 ☑  否 ☐" if v else "是 ☐  否 ☑"

        rdsj = 0 if data.get("rqls14", "") in ("群众", "") else 1

        ws["B3"] = data.get("rqls1", "")
        ws["I3"] = data.get("rqls2", "")
        ws["B6"] = yn(data.get("rqls3"))
        ws["B7"] = yn(data.get("rqls4"))
        ws["B8"] = yn(data.get("rqls5"), rdsj)
        ws["H6"] = yn(data.get("rqls6"))
        ws["H7"] = yn(data.get("rqls7"))
        ws["H8"] = yn(data.get("rqls8"), rdsj)
        ws["K6"] = yn(data.get("rqls9"))
        ws["K7"] = yn(data.get("rqls10"))
        ws["K8"] = yn(data.get("rqls11"))
        ws["N6"] = data.get("rqls12", "") or ""
        ws["N7"] = data.get("rqls13", "") or ""
        ws["N8"] = data.get("rqls14", "") or ""
        ws["D9"] = data.get("rqls15", "") or ""
        ws["D10"] = data.get("rqls16", "") or ""
        ws["L9"] = data.get("rqls17", "") or ""
        ws["L10"] = data.get("rqls18", "") or ""
        ws["B11"] = data.get("rqls19", "") or ""
        ws["I11"] = data.get("rqls20", "") or ""
        ws["N11"] = data.get("rqls21", "") or ""
        ws["D12"] = yn(data.get("rqls22"))
        ws["L12"] = yn(data.get("rqls23"))
        ws["D13"] = data.get("rqls24", "") or ""
        ws["L13"] = data.get("rqls25", "") or ""
        ws["D14"] = data.get("rqls26", "") or ""
        ws["D15"] = data.get("rqls27", "") or ""
        ws["C17"] = data.get("rqls28", "") or ""
        ws["M17"] = data.get("rqls29", "") or ""

        for mr in merged:
            ws.merge_cells(str(mr))

        fn = f"{data.get('rqls1', '')}_{rsid}_任前联审.xlsx"
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        encoded_fn = quote(fn.encode("utf-8"))
        response["Content-Disposition"] = (
            f"attachment; filename=\"{encoded_fn}\"; filename*=UTF-8''{encoded_fn}"
        )
        return response
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()
