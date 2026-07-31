"""
档案目录打印 - PDF + Excel
"""

import logging
import os
import openpyxl
from io import BytesIO
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.conf import settings
from main.utils import _get_conn
from main.utils.decorators import archive_perm_required
from weasyprint import HTML
from openpyxl.styles import Border, Side, Alignment
from urllib.parse import quote

logger = logging.getLogger(__name__)

CATEGORIES = [
    ("一", "履历材料", None),
    ("二", "自传材料", None),
    ("三", "鉴定、考核材料", None),
    (
        "四",
        "学历学位、职称、学术、培训等材料",
        [
            ("4-1", "学历学位材料"),
            ("4-2", "专业技术职务材料"),
            ("4-3", "科研学术材料"),
            ("4-4", "培训材料"),
        ],
    ),
    ("五", "政审材料", None),
    ("六", "党团材料", None),
    ("七", "奖励材料", None),
    ("八", "处分材料", None),
    (
        "九",
        "工资、任免、出国、会议等材料",
        [
            ("9-1", "工资材料"),
            ("9-2", "任免材料"),
            ("9-3", "出国材料"),
            ("9-4", "会议代表材料"),
        ],
    ),
    ("十", "其他材料", None),
]

# 四和九是父类，不插入空行
NO_EMPTY_CATS = {"四", "九"}

TEMPLATE_PATH = os.path.join(
    settings.BASE_DIR, "static", "excel_templates", "干部档案目录.xlsx"
)


def _load_data(rsid):
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT XM FROM RS_INFO WHERE RSID=?", (int(rsid),))
        person_name = (cursor.fetchone() or [""])[0]
        cursor.execute(
            """SELECT JBBH+'-'+CAST(XH AS VARCHAR(5)) AS 编号, CLTM AS 材料名称,
                   FYEAR AS 年, FMONTH AS 月, FDAY AS 日, YS AS 页数, BZ AS 备注
            FROM RS_ARCHINFO a LEFT JOIN CATETREE b ON a.FL=b.FL
            WHERE a.RSID=? ORDER BY a.FL, a.XH""",
            (int(rsid),),
        )
        rows = [
            dict(zip([c[0] for c in cursor.description], r)) for r in cursor.fetchall()
        ]
        cursor.close()
        data_map = {}
        for row in rows:
            parts = (row["编号"] or "").split("-")
            key = (
                "-".join(parts[:2])
                if len(parts) >= 3 and parts[0].isdigit()
                else parts[0]
            )
            data_map.setdefault(key, []).append(row)
        return person_name, data_map
    finally:
        if conn:
            conn.close()


@login_required
@archive_perm_required
def print_page(request):
    return render(
        request, "archives/directory_print.html", {"rsid": request.GET.get("rsid", "")}
    )


@login_required
def print_pdf_api(request):
    rsid = request.GET.get("rsid", "")
    empty_rows = int(request.GET.get("emptyRows", 5))
    if not rsid:
        return HttpResponse("缺少 rsid", status=400)

    person_name, data_map = _load_data(rsid)

    def font_size(text):
        l = len(text)
        if l > 25:
            return "10pt"
        if l > 18:
            return "12pt"
        return "14pt"

    def empty_rows_html(n):
        return (
            '<tr class="er"><td class="ca"></td><td class="cb"></td><td class="cc"></td><td class="cd"></td><td class="ce"></td><td class="cf"></td><td class="cg"></td></tr>'
            * n
        )

    def rows_html(cat_key, data_key=None):
        items = data_map.get(data_key or cat_key, [])
        h = ""
        for i, item in enumerate(items, 1):
            mat = item.get("材料名称") or ""
            bz = item.get("备注") or ""
            h += f"<tr class='dr'><td class='ca'>{i}</td><td class='cb' style='font-size:{font_size(mat)}'>{mat}</td><td class='cc'>{item.get('年') or ''}</td><td class='cd'>{item.get('月') or ''}</td><td class='ce'>{item.get('日') or ''}</td><td class='cf'>{item.get('页数') or ''}</td><td class='cg' style='font-size:{font_size(bz)}'>{bz}</td></tr>"
        # 父类不加空行
        if cat_key not in NO_EMPTY_CATS:
            h += empty_rows_html(empty_rows)
        return h

    cats = [
        ("一", "履历材料", True, "一"),
        ("二", "自传材料", True, "二"),
        ("三", "鉴定、考核材料", True, "三"),
        ("四", "学历学位、职称、学术、培训等材料", True, None),
        ("4-1", "学历学位材料", False, "4-1"),
        ("4-2", "专业技术职务材料", False, "4-2"),
        ("4-3", "科研学术材料", False, "4-3"),
        ("4-4", "培训材料", False, "4-4"),
        ("五", "政审材料", True, "五"),
        ("六", "党团材料", True, "六"),
        ("七", "奖励材料", True, "七"),
        ("八", "处分材料", True, "八"),
        ("九", "工资、任免、出国、会议等材料", True, None),
        ("9-1", "工资材料", False, "9-1"),
        ("9-2", "任免材料", False, "9-2"),
        ("9-3", "出国材料", False, "9-3"),
        ("9-4", "会议代表材料", False, "9-4"),
        ("十", "其他材料", True, "十"),
    ]

    body = ""
    for cat_key, cat_name, is_main, data_key in cats:
        cls = "cr" if is_main else "scr"
        body += f"<tr class='{cls}'><td class='ca'>{cat_key}</td><td class='cb' colspan='6'>{cat_name}</td></tr>"
        body += rows_html(cat_key, data_key)

    html = f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8"><style>
@page{{size:A4 portrait;margin:{1.5 / 2.54 * 72}pt {1.8 / 2.54 * 72}pt {2.0 / 2.54 * 72}pt {1.6 / 2.54 * 72}pt;@bottom-center{{content:"共 " counter(pages) " 页  第 " counter(page) " 页";font-size:9pt;font-family:"SimSun",sans-serif;}}}}
body{{font-family:"SimSun","宋体",sans-serif;margin:0;padding:0;font-size:14pt;}}
table{{border-collapse:collapse;width:100%;table-layout:fixed;}}
.col-ca{{width:6%;}}.col-cb{{width:53%;}}.col-cc{{width:10%;}}.col-cd{{width:4%;}}.col-ce{{width:4%;}}.col-cf{{width:5%;}}.col-cg{{width:12%;}}
th,td{{border:0.5px solid #000;padding:3px 5px;line-height:1.3;}}
.ca{{text-align:center;}}.cb{{text-align:left;}}.cc{{text-align:center;}}.cd{{text-align:center;}}.ce{{text-align:center;}}.cf{{text-align:center;}}.cg{{text-align:center;}}
.title{{text-align:center;font-size:22pt;font-weight:bold;font-family:"SimHei","黑体",sans-serif;margin-bottom:6pt;}}
.name-line{{font-size:16pt;font-family:"SimSun","宋体",sans-serif;margin-bottom:4pt;}}
.cr td{{font-weight:bold;font-family:"SimHei","黑体",sans-serif;font-size:16pt;text-align:left;padding:4px 6px;}}
.scr td{{font-family:"SimSun","宋体",sans-serif;font-size:14pt;text-align:left;padding:3px 6px;}}
.dr td{{text-align:center;font-size:14pt;}}.dr td:nth-child(2){{text-align:left;}}
.er td{{height:24pt;padding:0;border:0.5px solid #000;}}
thead{{display:table-header-group;}}
.hr th{{text-align:center;font-weight:bold;font-family:"SimHei","黑体",sans-serif;font-size:14pt;}}
</style></head><body>
<div class="title">干部档案目录</div><div class="name-line">姓名：{person_name}</div>
<table>
<colgroup>
<col class="col-ca"><col class="col-cb"><col class="col-cc"><col class="col-cd"><col class="col-ce"><col class="col-cf"><col class="col-cg">
</colgroup>
<thead>
<tr class="hr"><th class="ca" rowspan="2">序号</th><th class="cb" rowspan="2">材 料 题 名</th><th class="cc" colspan="3">材料形成时间</th><th class="cf" rowspan="2">页数</th><th class="cg" rowspan="2">备注</th></tr>
<tr class="hr"><th class="cc">年</th><th class="cd">月</th><th class="ce">日</th></tr>
</thead><tbody>{body}</tbody></table></body></html>"""
    buf = BytesIO()
    HTML(string=html).write_pdf(buf)
    buf.seek(0)
    resp = HttpResponse(buf, content_type="application/pdf")
    resp["Content-Security-Policy"] = "frame-ancestors 'self'"
    resp["X-Frame-Options"] = "SAMEORIGIN"
    return resp


@login_required
def print_export_api(request):
    rsid = request.GET.get("rsid", "")
    empty_rows = int(request.GET.get("emptyRows", 5))
    if not rsid:
        return HttpResponse("缺少 rsid", status=400)

    person_name, data_map = _load_data(rsid)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "干部档案目录"

    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "portrait"
    ws.print_title_rows = "$3:$4"
    ws.oddFooter.center.text = "第 &P 页，共 &N 页"
    ws.evenFooter.center.text = "第 &P 页，共 &N 页"
    ws.page_margins.left = 1.6 / 2.54
    ws.page_margins.right = 1.8 / 2.54
    ws.page_margins.top = 1.5 / 2.54
    ws.page_margins.bottom = 2.0 / 2.54

    col_widths = [5, 50, 8, 5, 5, 5, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    title_font = openpyxl.styles.Font(name="黑体", size=22, bold=True)
    name_font = openpyxl.styles.Font(name="宋体", size=16)
    header_font = openpyxl.styles.Font(name="黑体", size=16, bold=True)
    cat_font = openpyxl.styles.Font(name="黑体", size=16, bold=True)
    sub_font = openpyxl.styles.Font(name="宋体", size=16)
    data_font = openpyxl.styles.Font(name="宋体", size=16)
    shrink_font = openpyxl.styles.Font(name="宋体", size=16)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    shrink_align = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)

    def sc(r, c, v, font=data_font, align=center_align):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = font
        cell.alignment = align
        return cell

    ws.merge_cells("A1:G1")
    sc(1, 1, "干部档案目录", title_font)
    ws.merge_cells("A2:C2")
    sc(2, 1, f"姓名：{person_name or ''}", name_font, left_align)

    ws.merge_cells("A3:A4")
    sc(3, 1, "序号", header_font)
    ws.merge_cells("B3:B4")
    sc(3, 2, "材  料  题  名", header_font)
    ws.merge_cells("C3:E3")
    sc(3, 3, "材料形成时间", header_font)
    sc(4, 3, "年", header_font)
    sc(4, 4, "月", header_font)
    sc(4, 5, "日", header_font)
    ws.merge_cells("F3:F4")
    sc(3, 6, "页数", header_font)
    ws.merge_cells("G3:G4")
    sc(3, 7, "备注", header_font)

    row = 5
    cats = [
        ("一", "履历材料", True, "一"),
        ("二", "自传材料", True, "二"),
        ("三", "鉴定、考核材料", True, "三"),
        ("四", "学历学位、职称、学术、培训等材料", True, None),
        ("4-1", "学历学位材料", False, "4-1"),
        ("4-2", "专业技术职务材料", False, "4-2"),
        ("4-3", "科研学术材料", False, "4-3"),
        ("4-4", "培训材料", False, "4-4"),
        ("五", "政审材料", True, "五"),
        ("六", "党团材料", True, "六"),
        ("七", "奖励材料", True, "七"),
        ("八", "处分材料", True, "八"),
        ("九", "工资、任免、出国、会议等材料", True, None),
        ("9-1", "工资材料", False, "9-1"),
        ("9-2", "任免材料", False, "9-2"),
        ("9-3", "出国材料", False, "9-3"),
        ("9-4", "会议代表材料", False, "9-4"),
        ("十", "其他材料", True, "十"),
    ]

    for cat_key, cat_name, is_main, data_key in cats:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        font = cat_font if is_main else sub_font
        sc(row, 1, f"{cat_key}、{cat_name}", font, left_align)
        ws.row_dimensions[row].height = 24
        row += 1

        items = data_map.get(data_key or cat_key, [])
        for j, item in enumerate(items):
            ws.row_dimensions[row].height = 24
            sc(row, 1, j + 1)
            sc(row, 2, item.get("材料名称", ""), shrink_font, shrink_align)
            sc(row, 3, item.get("年", ""))
            sc(row, 4, item.get("月", ""))
            sc(row, 5, item.get("日", ""))
            sc(row, 6, item.get("页数", ""))
            sc(row, 7, item.get("备注", ""), shrink_font, shrink_align)
            row += 1

        if cat_key not in NO_EMPTY_CATS:
            for _ in range(empty_rows):
                ws.row_dimensions[row].height = 24
                for c in range(1, 8):
                    sc(row, c, "")
                row += 1

    for r in range(3, row):
        for c in range(1, 8):
            ws.cell(row=r, column=c).border = thin_border

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{person_name}_干部档案目录.xlsx"
    response = HttpResponse(
        buf,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    encoded_fn = quote(filename.encode("utf-8"))
    response["Content-Disposition"] = (
        f"attachment; filename=\"{encoded_fn}\"; filename*=UTF-8''{encoded_fn}"
    )
    return response
