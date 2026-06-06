"""专项审核情况登记表"""

import hashlib
import json
import logging
import os
from io import BytesIO
from urllib.parse import quote

import openpyxl
from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.csrf import csrf_exempt

from main.db_utils import _get_conn
from main.decorators import archive_perm_required

logger = logging.getLogger(__name__)


def _get_yhbh(request):
    try:
        return request.user.profile.yhbh or 0
    except Exception:
        return 0


@login_required
@archive_perm_required
def person_audit_view(request):
    return render(request, "archives/person_audit.html")


@login_required
@archive_perm_required
def audit_data_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        yhbh = _get_yhbh(request)
        cursor.execute("{CALL ZXSH_DAZSQKDJB_WEB(?, 'SELECT', ?)}", (int(rsid), yhbh))
        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        data = dict(zip(cols, row)) if row else {}
        cursor.close()
        return JsonResponse({"code": 0, "data": data})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@csrf_exempt
def audit_save_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({"code": 400})

    rsid = data.get("rsid")
    if not rsid:
        return JsonResponse({"code": 400, "msg": "缺少RSID"})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        yhbh = _get_yhbh(request)

        cursor.execute("{CALL ZXSH_DAZSQKDJB_WEB(?, 'SELECT', ?)}", (int(rsid), yhbh))
        cols = [col[0] for col in cursor.description]
        old_row = cursor.fetchone()
        old_data = dict(zip(cols, old_row)) if old_row else {}

        merged = {**old_data, **data}
        merged["RSID"] = int(rsid)
        merged["userid"] = yhbh

        fields = [
            "RSID",
            "types",
            "userid",
            "cssj1A",
            "cssj2A",
            "cssj3A",
            "cssj4A",
            "cssj5A",
            "cssj6A",
            "cssj7A",
            "cssj8A",
            "cjgz1A",
            "cjgz2A",
            "cjgz3A",
            "cjgz4A",
            "cjgz5A",
            "cjgz6A",
            "cjgz7A",
            "rdsj1A",
            "rdsj2A",
            "rdsj3A",
            "rdsj4A",
            "rdsj5A",
            "rdsj6A",
            "rdsj7A",
            "rdsj8A",
            "rdsj9A",
            "rdsj10A",
            "xlxw1A",
            "xlxw2A",
            "xlxw3A",
            "xlxw4A",
            "xlxw5A",
            "xlxw6A",
            "xlxw7A",
            "xlxw8A",
            "xlxw9A",
            "xlxw10A",
            "xlxw11A",
            "xlxw12A",
            "xlxw13A",
            "gzjl1A",
            "gzjl2A",
            "gzjl3A",
            "gzjl4A",
            "gzjl5A",
            "gzjl6A",
            "gzjl7A",
            "gbsf1A",
            "gbsf2A",
            "gbsf3A",
            "gbsf4A",
            "zyjs1A",
            "zyjs2A",
            "zyjs3A",
            "zyjs4A",
            "jcqk1A",
            "jcqk2A",
            "mc1A",
            "mc2A",
            "mc3A",
            "shgx1A",
            "shgx2A",
            "qtwt",
            "shyj",
            "cssj2B",
            "cssj3B",
            "cssj4B",
            "cssj5B",
            "cssj6B",
            "cssj7B",
            "cssj8B",
            "cjgz3B",
            "cjgz5B",
            "cjgz6B",
            "cjgz7B",
            "rdsj6B",
            "rdsj8B",
            "rdsj9B",
            "rdsj10B",
            "xlxw11B",
            "xlxw12B",
            "xlxw13B",
            "gzjl4B",
            "gzjl7B",
            "gbsf2B",
            "gbsf3B",
            "gbsf4B",
            "zyjs1B",
            "zyjs2B",
            "zyjs3B",
            "zyjs4B",
            "mc3B",
            "cssj2C",
            "cssj4C",
            "cssj5C",
            "cssj6C",
            "cssj7C",
            "cssj8C",
            "cjgz3C",
            "cjgz5C",
            "cjgz6C",
            "cjgz7C",
            "rdsj8C",
            "rdsj9C",
            "rdsj10C",
            "xlxw11C",
            "xlxw12C",
            "xlxw13C",
            "gzjl4C",
            "gzjl7C",
            "gbsf2C",
            "gbsf3C",
            "gbsf4C",
            "zyjs1C",
            "zyjs2C",
            "zyjs3C",
            "zyjs4C",
            "mc3C",
            "cssj2D",
            "cssj3D",
            "cssj4D",
            "cssj5D",
            "cssj6D",
            "cssj7D",
            "cssj8D",
            "cjgz3D",
            "cjgz5D",
            "cjgz6D",
            "cjgz7D",
            "rdsj6D",
            "rdsj8D",
            "rdsj9D",
            "rdsj10D",
            "xlxw11D",
            "xlxw12D",
            "xlxw13D",
            "gzjl4D",
            "gzjl7D",
            "gbsf2D",
            "gbsf3D",
            "gbsf4D",
            "zyjs1D",
            "zyjs2D",
            "zyjs3D",
            "mc3D",
            "cssj",
            "cjgz",
            "rdsj",
            "xlxw",
            "gzjl",
            "gbsf",
            "zyjs",
            "jcqk",
            "mc",
            "shgx",
            "shr1",
            "shr2",
            "shrtime",
            "shrtime2",
        ]

        values = []
        for f in fields:
            if f == "RSID":
                values.append(int(rsid))
            elif f == "types":
                values.append("UPDATE")
            elif f == "userid":
                values.append(yhbh)
            else:
                v = merged.get(f, "")
                if isinstance(v, bool):
                    v = 1 if v else 0
                values.append(v)

        placeholders = ",".join(["?"] * len(values))
        sql = f"{{CALL ZXSH_DAZSQKDJB_WEB({placeholders})}}"
        cursor.execute(sql, values)
        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "保存成功"})
    except Exception as e:
        logger.error(f"专审保存失败: {e}")
        if conn:
            try:
                conn.rollback()
            except:  # noqa: E722
                pass
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            try:
                conn.close()
            except:  # noqa: E722
                pass


@login_required
def audit_proof_check_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        yhbh = _get_yhbh(request)
        cursor.execute("{CALL ZXSH_RQLSDJB(?, 'SELECTZXSHZM', ?)}", (int(rsid), yhbh))

        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        result = dict(zip(cols, row)) if row else {}
        result = {k: ("" if isinstance(v, bytes) else v) for k, v in result.items()}

        details = []
        if cursor.nextset():
            dcols = [col[0] for col in cursor.description]
            for row in cursor.fetchall():
                detail = dict(zip(dcols, row))
                detail = {
                    k: ("" if isinstance(v, bytes) else v) for k, v in detail.items()
                }
                details.append(detail)

        cursor.close()
        return JsonResponse({"code": 0, "result": result, "details": details})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


def _generate_proof_pdf(rsid, proof_type, yhbh):
    import tempfile
    from datetime import datetime

    import qrcode
    from PIL import Image
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph

    font_path = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
    if not os.path.exists(font_path):
        font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    pdfmetrics.registerFont(TTFont("Chinese", font_path))

    conn = _get_conn()
    cursor = conn.cursor()
    cursor.execute("{CALL ZXSH_RQLSDJB(?, ?, ?)}", (int(rsid), proof_type, yhbh))
    cols = [col[0] for col in cursor.description]
    row = cursor.fetchone()
    data = dict(zip(cols, row)) if row else {}
    cursor.close()
    conn.close()

    photo_path = None
    if data.get("DQZP"):
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".jpg")
        tmp.write(data["DQZP"])
        tmp.close()
        photo_path = tmp.name

    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.ERROR_CORRECT_H,
        box_size=10,
        border=2,
    )
    qr.add_data(data.get("qrcode", ""))
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="black", back_color="white").convert("RGB")  # type: ignore

    if photo_path:
        logo = Image.open(photo_path).resize((80, 80))
        pos = (
            (qr_img.size[0] - logo.size[0]) // 2,
            (qr_img.size[1] - logo.size[1]) // 2,
        )
        qr_img.paste(logo, pos)

    qr_tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    qr_img.save(qr_tmp.name)
    qr_tmp.close()

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    c.setFont("Chinese", 22)
    c.drawCentredString(width / 2, height - 120, "干部人事档案专项审核证明")

    style = ParagraphStyle(
        "ChineseStyle",
        fontName="Chinese",
        fontSize=16,
        leading=32,  # 行间距，比字号大一些
        spaceBefore=6,  # 段前间距
        spaceAfter=6,  # 段后间距
        alignment=4,
    )
    text = data.get("paragraph", "").replace("\n", "<br/>")
    text = data.get("paragraph", "")
    text = text.replace(" ", "&nbsp;")
    text = text.replace("\n", "<br/>")
    p = Paragraph(text, style)
    p.wrapOn(c, width - 160, height - 160)
    p.drawOn(c, 80, height - 440)

    c.setFont("Chinese", 16)
    today = datetime.now().strftime("%Y年%m月%d日")
    c.drawRightString(width - 120, 200, "盘州市教育局档案室")
    c.drawRightString(width - 120, 165, today)

    c.drawImage(qr_tmp.name, 80, 150, width=200, height=200)
    c.showPage()
    c.save()
    buf.seek(0)

    if photo_path:
        os.unlink(photo_path)
    os.unlink(qr_tmp.name)

    return buf


@login_required
def audit_proof_print_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    try:
        yhbh = _get_yhbh(request)
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute("{CALL ZXSH_RQLSDJB(?, 'SELECTZXSHZM', ?)}", (int(rsid), yhbh))
        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        check_data = dict(zip(cols, row)) if row else {}
        cursor.close()
        conn.close()

        if check_data.get("ZSZM") != "yes":
            return JsonResponse(
                {"code": 400, "msg": "该同志档案未专审通过，不能出具证明"}
            )

        buf = _generate_proof_pdf(rsid, "SELECTZXSHZM", yhbh)
        resp = HttpResponse(buf.read(), content_type="application/pdf")
        resp["Content-Disposition"] = 'inline; filename="proof.pdf"'
        return resp
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})


@login_required
def audit_proof_print_admin_api(request):
    rsid = request.GET.get("r", "")
    pwd = request.GET.get("k", "")

    if hashlib.md5(pwd.encode()).hexdigest() != "f4d0e2e7fc057a58f7ca4a391f01940a":
        return JsonResponse({"code": 403})
    if not rsid:
        return JsonResponse({"code": 400})

    try:
        yhbh = _get_yhbh(request)
        buf = _generate_proof_pdf(rsid, "SELECTZXSHZM_Admin", yhbh)
        resp = HttpResponse(buf.read(), content_type="application/pdf")
        resp["Content-Disposition"] = 'inline; filename="proof.pdf"'
        return resp
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})


@login_required
@archive_perm_required
def audit_export_api(request):
    """导出专审登记表 Excel"""
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        yhbh = _get_yhbh(request)
        cursor.execute("{CALL ZXSH_DAZSQKDJB_WEB(?, 'SELECT', ?)}", (int(rsid), yhbh))
        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        data = dict(zip(cols, row)) if row else {}

        cursor.execute("SELECT XM FROM RS_INFO WHERE RSID=?", (int(rsid),))
        name_row = cursor.fetchone()
        name = name_row[0] if name_row else ""
        cursor.close()
    finally:
        if conn:
            conn.close()

    template_path = os.path.join(
        settings.BASE_DIR, "static", "excel_templates", "专项审核情况登记表.xlsx"
    )
    wb = openpyxl.load_workbook(template_path)
    ws = wb["专项审核情况登记表"]

    ws["A3"] = f"姓名：{name}"
    ws["E3"] = name
    ws["H3"] = "级别："

    def yn(v):
        if v is None:
            return ""
        return "是" if v else "否"

    def ynn(v):
        if v is None:
            return "□是  □否"
        return f"□{yn(v)}  □{yn(not v)}"

    def gv(k, d=""):
        v = data.get(k)
        return str(v) if v is not None else str(d)

    def textToTime(value, years):
        if not value:
            return "    年  月  日" if years else ""
        value = str(value)
        if len(value) == 6:
            return f"{value[:4]}年{int(value[4:6])}月"
        elif len(value) == 8:
            return f"{value[:4]}年{int(value[4:6])}月{int(value[6:8])}日"
        return value

    # ========== 出生年月 (C# R5-12 → openpyxl R6-13) ==========
    ws["E6"] = ynn(data.get("cssj1A"))
    ws["E7"] = textToTime(gv("cssj2A"), False)
    ws["E8"] = textToTime(gv("cssj3A"), False)
    ws["E9"] = ynn(data.get("cssj4A"))
    ws["E10"] = ynn(data.get("cssj5A"))
    ws["E11"] = ynn(data.get("cssj6A"))
    ws["E12"] = ynn(data.get("cssj7A"))
    ws["E13"] = ynn(data.get("cssj8A"))
    ws["F7"] = gv("cssj2B")
    ws["F8"] = gv("cssj3B")
    ws["F9"] = gv("cssj4B")
    ws["F10"] = gv("cssj5B")
    ws["F11"] = gv("cssj6B")
    ws["F12"] = gv("cssj7B")
    ws["F13"] = gv("cssj8B")
    ws["G7"] = gv("cssj2C")
    ws["G9"] = gv("cssj4C")
    ws["G10"] = gv("cssj5C")
    ws["G11"] = gv("cssj6C")
    ws["G12"] = gv("cssj7C")
    ws["G13"] = gv("cssj8C")
    ws["H7"] = textToTime(gv("cssj2D"), True)
    ws["H8"] = textToTime(gv("cssj3D"), True)
    ws["H9"] = textToTime(gv("cssj4D"), True)
    ws["H10"] = textToTime(gv("cssj5D"), True)
    ws["H11"] = textToTime(gv("cssj6D"), True)
    ws["H12"] = textToTime(gv("cssj7D"), True)
    ws["H13"] = textToTime(gv("cssj8D"), True)
    ws["I6"] = gv("cssj")

    # ========== 参加工作时间 (C# R13-19 → openpyxl R14-20) ==========
    ws["E14"] = ynn(data.get("cjgz1A"))
    ws["E15"] = ynn(data.get("cjgz2A"))
    ws["E16"] = textToTime(gv("cjgz3A"), False)
    ws["E17"] = textToTime(gv("cjgz4A"), False)
    ws["E18"] = ynn(data.get("cjgz5A"))
    ws["E19"] = ynn(data.get("cjgz6A"))
    ws["E20"] = ynn(data.get("cjgz7A"))
    ws["F16"] = gv("cjgz3B")
    ws["F18"] = gv("cjgz5B")
    ws["F19"] = gv("cjgz6B")
    ws["F20"] = gv("cjgz7B")
    ws["G16"] = gv("cjgz3C")
    ws["G18"] = gv("cjgz5C")
    ws["G19"] = gv("cjgz6C")
    ws["G20"] = gv("cjgz7C")
    ws["H16"] = textToTime(gv("cjgz3D"), True)
    ws["H18"] = textToTime(gv("cjgz5D"), True)
    ws["H19"] = textToTime(gv("cjgz6D"), True)
    ws["H20"] = textToTime(gv("cjgz7D"), True)
    ws["I14"] = gv("cjgz")

    # ========== 入党时间 (C# R22-31 → openpyxl R23-32) ==========
    ws["E23"] = ynn(data.get("rdsj1A"))
    ws["E24"] = ynn(data.get("rdsj2A"))
    ws["E25"] = ynn(data.get("rdsj3A"))
    ws["E26"] = ynn(data.get("rdsj4A"))
    ws["E27"] = ynn(data.get("rdsj5A"))
    ws["E28"] = textToTime(gv("rdsj6A"), False)
    ws["E29"] = textToTime(gv("rdsj7A"), False)
    ws["E30"] = ynn(data.get("rdsj8A"))
    ws["E31"] = ynn(data.get("rdsj9A"))
    ws["E32"] = ynn(data.get("rdsj10A"))
    ws["F28"] = gv("rdsj6B")
    ws["F30"] = gv("rdsj8B")
    ws["F31"] = gv("rdsj9B")
    ws["F32"] = gv("rdsj10B")
    ws["G30"] = gv("rdsj8C")
    ws["G31"] = gv("rdsj9C")
    ws["G32"] = gv("rdsj10C")
    ws["H28"] = textToTime(gv("rdsj6D"), True)
    ws["H30"] = textToTime(gv("rdsj8D"), True)
    ws["H31"] = textToTime(gv("rdsj9D"), True)
    ws["H32"] = textToTime(gv("rdsj10D"), True)
    ws["I23"] = gv("rdsj")

    # ========== 学历学位 (C# R32-44 → openpyxl R33-45) ==========
    ws["E33"] = ynn(data.get("xlxw1A"))
    ws["E34"] = ynn(data.get("xlxw2A"))
    ws["E35"] = gv("xlxw3A")
    ws["E36"] = gv("xlxw4A")
    ws["E37"] = gv("xlxw5A")
    ws["E38"] = gv("xlxw6A")
    ws["E39"] = ynn(data.get("xlxw7A"))
    ws["E40"] = ynn(data.get("xlxw8A"))
    ws["E41"] = ynn(data.get("xlxw9A"))
    ws["E42"] = ynn(data.get("xlxw10A"))
    ws["E43"] = ynn(data.get("xlxw11A"))
    ws["E44"] = ynn(data.get("xlxw12A"))
    ws["E45"] = ynn(data.get("xlxw13A"))
    ws["F43"] = gv("xlxw11B")
    ws["F44"] = gv("xlxw12B")
    ws["F45"] = gv("xlxw13B")
    ws["G43"] = gv("xlxw11C")
    ws["G44"] = gv("xlxw12C")
    ws["G45"] = gv("xlxw13C")
    ws["H43"] = textToTime(gv("xlxw11D"), True)
    ws["H44"] = textToTime(gv("xlxw12D"), True)
    ws["H45"] = textToTime(gv("xlxw13D"), True)
    ws["I33"] = gv("xlxw")

    # ========== 工作经历 (C# R47-53 → openpyxl R48-54) ==========
    ws["E48"] = ynn(data.get("gzjl1A"))
    ws["E49"] = ynn(data.get("gzjl2A"))
    ws["E50"] = ynn(data.get("gzjl3A"))
    ws["E51"] = ynn(data.get("gzjl4A"))
    ws["E52"] = ynn(data.get("gzjl5A"))
    ws["E53"] = ynn(data.get("gzjl6A"))
    ws["E54"] = ynn(data.get("gzjl7A"))
    ws["F51"] = gv("gzjl4B")
    ws["F54"] = gv("gzjl7B")
    ws["G51"] = gv("gzjl4C")
    ws["G54"] = gv("gzjl7C")
    ws["H51"] = textToTime(gv("gzjl4D"), True)
    ws["H54"] = textToTime(gv("gzjl7D"), True)
    ws["I48"] = gv("gzjl")

    # ========== 干部身份 (C# R54-57 → openpyxl R55-58) ==========
    ws["E55"] = ynn(data.get("gbsf1A"))
    ws["E56"] = ynn(data.get("gbsf2A"))
    ws["E57"] = ynn(data.get("gbsf3A"))
    ws["E58"] = ynn(data.get("gbsf4A"))
    ws["F56"] = gv("gbsf2B")
    ws["F57"] = gv("gbsf3B")
    ws["F58"] = gv("gbsf4B")
    ws["G56"] = gv("gbsf2C")
    ws["G57"] = gv("gbsf3C")
    ws["G58"] = gv("gbsf4C")
    ws["H56"] = textToTime(gv("gbsf2D"), True)
    ws["H57"] = textToTime(gv("gbsf3D"), True)
    ws["H58"] = textToTime(gv("gbsf4D"), True)
    ws["I55"] = gv("gbsf")

    # ========== 专业技术职务 (C# R58-61 → openpyxl R59-62) ==========
    ws["E59"] = gv("zyjs1A")
    ws["E60"] = ynn(data.get("zyjs2A"))
    ws["E61"] = ynn(data.get("zyjs3A"))
    ws["E62"] = ynn(data.get("zyjs4A"))
    ws["F59"] = gv("zyjs1B")
    ws["F60"] = gv("zyjs2B")
    ws["F61"] = gv("zyjs3B")
    ws["F62"] = gv("zyjs4B")
    ws["G59"] = gv("zyjs1C")
    ws["G60"] = gv("zyjs2C")
    ws["G61"] = gv("zyjs3C")
    ws["G62"] = gv("zyjs4C")
    ws["H59"] = textToTime(gv("zyjs1D"), True)
    ws["H60"] = textToTime(gv("zyjs2D"), True)
    ws["H61"] = textToTime(gv("zyjs3D"), True)
    ws["I59"] = gv("zyjs")

    # ========== 奖惩情况 (C# R62-63 → openpyxl R63-64) ==========
    ws["E63"] = ynn(data.get("jcqk1A"))
    ws["E64"] = ynn(data.get("jcqk2A"))
    ws["I63"] = gv("jcqk")

    # ========== 民族 (C# R64-66 → openpyxl R65-67) ==========
    ws["E65"] = gv("mc1A")
    ws["E66"] = ynn(data.get("mc2A"))
    ws["E67"] = ynn(data.get("mc3A"))
    ws["F67"] = gv("mc3B")
    ws["G67"] = gv("mc3C")
    ws["H67"] = textToTime(gv("mc3D"), True)
    ws["I65"] = gv("mc")

    # ========== 家庭主要成员 (C# R69-70 → openpyxl R70-71) ==========
    ws["E70"] = ynn(data.get("shgx1A"))
    ws["E71"] = ynn(data.get("shgx2A"))
    ws["I70"] = gv("shgx")

    # ========== 其他 (C# R71-72 → openpyxl R72-73) ==========
    ws["B72"] = gv("qtwt")
    ws["B73"] = gv("shyj")

    # ========== 审核人 (C# R73 → openpyxl R74) ==========
    shr1 = gv("shr1")
    shr2 = gv("shr2")
    shrtime = textToTime(gv("shrtime"), True)
    ws["A74"] = (
        f"初审人：{shr1}         初审时间：{shrtime}         复审人：{shr2}     复审时间：{shrtime}"
    )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = (
        f"attachment; filename={quote(name + '_专项审核情况登记表.xlsx')}"
    )
    return resp
