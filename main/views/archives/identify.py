#'''认定表'''
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.conf import settings
import json
import logging
import os
import re
from io import BytesIO
from main.utils import _get_conn
import openpyxl
from urllib.parse import quote

from main.utils.decorators import archive_perm_required

# @archive_perm_required
logger = logging.getLogger(__name__)


def _get_yhbh(request):
    try:
        return request.user.profile.yhbh or 0
    except:
        return 0


@login_required
@archive_perm_required
def person_identify_view(request):
    return render(request, "archives/person_identify.html")


# ==================== 认定表数据 ====================


@login_required
def identify_data_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT a.cssj2C AS RDBa2, a.cssj AS RDBa3,
                   a.cjgz3C AS RDBb2, a.cjgz AS RDBb3,
                   a.rdsj AS RDBc3,
                   a.rdyj AS RDBd1, a.zzbmyj AS RDBd2,
                   b.CSNY AS RDBa1, b.WORKTIME AS RDBb1, b.JOINTIME AS RDBc1,
                   b.DCDW AS RDBe1, b.DCSJ AS RDBe2, b.DJDW AS RDBe3,
                   b.DJSJ AS RDBe4, b.AR AS RDBe5, b.WYZ AS RDBe6, b.DCYY AS RDBe7,
                   b.XM AS xm, b.JOBUNIT
            FROM YW_ZXSHDJ a
            LEFT JOIN RS_INFO b ON a.RSID = b.RSID
            WHERE a.RSID = ?
        """,
            (int(rsid),),
        )
        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        data = dict(zip(cols, row)) if row else {}

        if data.get("RDBc1") and not data.get("RDBc2"):
            data["RDBc2"] = "入党志愿书"

        # 拆分 RDBd1（用 \r\n卍 分隔，最少1个卍，最多2个卍）
        rdbd1 = data.get("RDBd1") or ""
        parts = re.split(r"\r?\n?卍", rdbd1)
        if len(parts) == 2:
            data["RDBd1_main"] = parts[0]
            data["JL01"] = ""
            data["JL02"] = parts[1]
        elif len(parts) >= 3:
            data["RDBd1_main"] = parts[0]
            data["JL01"] = parts[1]
            data["JL02"] = parts[2]
        else:
            data["RDBd1_main"] = rdbd1
            data["JL01"] = ""
            data["JL02"] = ""

        # 涂改记录
        cursor.execute(
            """
            SELECT xh AS 序号, lx AS 涂改类型, name AS 被涂改材料名称,
                   createTime AS 被涂改材料形成时间, totalTime AS 涂改后时间
            FROM YW_Alteration WHERE rsid=? ORDER BY xh
        """,
            (int(rsid),),
        )
        alt_cols = [col[0] for col in cursor.description]
        alt_rows = [dict(zip(alt_cols, r)) for r in cursor.fetchall()]
        cursor.close()

        return JsonResponse({"code": 0, "data": data, "alterations": alt_rows})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


# ==================== 保存认定表 ====================


@login_required
@csrf_exempt
def identify_save_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except:
        return JsonResponse({"code": 400})

    rsid = data.get("rsid")
    if not rsid:
        return JsonResponse({"code": 400, "msg": "缺少RSID"})

    rdbd1 = data.get("RDBd1_main", "") or ""
    jl01 = data.get("JL01", "") or ""
    jl02 = data.get("JL02", "") or ""

    if jl01:
        full_rdbd1 = f"{rdbd1}\r\n卍{jl01}\r\n卍{jl02}"
    else:
        full_rdbd1 = f"{rdbd1}\r\n卍{jl02}"

    rdbd2 = data.get("RDBd2", "") or ""
    yhbh = _get_yhbh(request)

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()

        cursor.execute(
            """
            UPDATE YW_ZXSHDJ SET
                cssj2C = ?, cssj = ?,
                cjgz3C = ?, cjgz = ?,
                rdsj = ?,
                rdyj = ?, zzbmyj = ?
            WHERE RSID = ?
        """,
            (
                data.get("RDBa2", ""),
                data.get("RDBa3", ""),
                data.get("RDBb2", ""),
                data.get("RDBb3", ""),
                data.get("RDBc3", ""),
                full_rdbd1,
                rdbd2,
                int(rsid),
            ),
        )

        # 日志
        cursor.execute("SELECT CSNY FROM RS_INFO WHERE RSID=?", (int(rsid),))
        old_csny = (cursor.fetchone() or [None])[0]
        new_csny = data.get("RDBa1", "")
        if str(old_csny or "") != str(new_csny):
            cursor.execute(
                "INSERT INTO Z_Log_rsinfo (userid, rsid, OperationType, oldvalue, newvalue, classNumber, mtime) VALUES (?, ?, 1, ?, ?, '出生年月', GETDATE())",
                (yhbh, int(rsid), str(old_csny or ""), str(new_csny)),
            )

        cursor.execute("SELECT WORKTIME FROM RS_INFO WHERE RSID=?", (int(rsid),))
        old_wt = (cursor.fetchone() or [None])[0]
        new_wt = data.get("RDBb1", "")
        if str(old_wt or "") != str(new_wt):
            cursor.execute(
                "INSERT INTO Z_Log_rsinfo (userid, rsid, OperationType, oldvalue, newvalue, classNumber, mtime) VALUES (?, ?, 1, ?, ?, '参工时间', GETDATE())",
                (yhbh, int(rsid), str(old_wt or ""), str(new_wt)),
            )

        cursor.execute(
            """
            UPDATE RS_INFO SET
                CSNY = ?, WORKTIME = ?, JOINTIME = ?,
                DCDW = ?, DCSJ = ?, DJDW = ?, DJSJ = ?,
                AR = ?, WYZ = ?, DCYY = ?
            WHERE RSID = ?
        """,
            (
                data.get("RDBa1", ""),
                data.get("RDBb1", ""),
                data.get("RDBc1", ""),
                data.get("RDBe1", ""),
                data.get("RDBe2", ""),
                data.get("RDBe3", ""),
                data.get("RDBe4", ""),
                data.get("RDBe5", ""),
                data.get("RDBe6", ""),
                data.get("RDBe7", ""),
                int(rsid),
            ),
        )

        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "保存成功"})
    except Exception as e:
        logger.error(f"认定表保存失败: {e}")
        if conn:
            try:
                conn.rollback()
            except:
                pass
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass


# ==================== 生成认定意见 ====================


@login_required
def identify_generate_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})
    jl01_from_page = request.GET.get("JL01", "")

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT a.cssj2C AS RDBa2, a.cssj AS RDBa3,
                   a.cjgz3C AS RDBb2, a.cjgz AS RDBb3,
                   b.CSNY AS RDBa1, b.WORKTIME AS RDBb1, b.AR AS RDBe5,
                   b.XM AS xm
            FROM YW_ZXSHDJ a
            LEFT JOIN RS_INFO b ON a.RSID = b.RSID
            WHERE a.RSID = ?
        """,
            (int(rsid),),
        )
        cols = [col[0] for col in cursor.description]
        row = cursor.fetchone()
        rdb = dict(zip(cols, row)) if row else {}

        if not rdb:
            cursor.close()
            conn.close()
            return JsonResponse({"code": 400, "msg": "请先填写专审情况登记表"})

        cursor.execute(
            "SELECT lx, name, createTime, totalTime FROM YW_Alteration WHERE rsid=?",
            (int(rsid),),
        )
        alt_rows = cursor.fetchall()
        cursor.close()
        conn.close()

        def fmt6(s):
            if not s:
                return ""
            s = str(s)
            if len(s) >= 6:
                return s[:4] + "年" + str(int(s[4:6])) + "月"
            return s

        def fmt8(s):
            if not s:
                return ""
            s = str(s)
            if len(s) >= 8:
                return s[:4] + "年" + str(int(s[4:6])) + "月" + str(int(s[6:8])) + "日"
            return s

        def fmt_long(s):
            if not s:
                return ""
            s = str(s)
            times = [t.strip() for t in s.replace("、", ",").split(",") if t.strip()]
            result = []
            for t in times:
                if len(t) == 6:
                    result.append(t[:4] + "年" + str(int(t[4:6])) + "月")
                elif len(t) == 4:
                    result.append(t + "年")
                else:
                    result.append(t)
            return "、".join(result)

        a1 = fmt6(rdb.get("RDBa1", ""))
        a2 = rdb.get("RDBa2", "")
        a3 = fmt_long(rdb.get("RDBa3", ""))
        b1 = fmt6(rdb.get("RDBb1", ""))
        b2 = rdb.get("RDBb2", "")
        b3 = fmt_long(rdb.get("RDBb3", ""))
        e5 = fmt8(rdb.get("RDBe5", ""))
        name = rdb.get("xm", "")

        csny_alt = ""
        cjgz_alt = ""
        for alt in alt_rows:
            lx = alt[0] or ""
            mat = alt[1] or ""
            t1 = fmt8(alt[2])
            t2 = fmt6(alt[3])
            desc = f"{t1}形成的《{mat}》中有涂改的情况，涂改后"
            if lx == "出生年月":
                csny_alt += desc + f"出生年月为{t2}。"
            elif lx == "参加工作时间":
                cjgz_alt += desc + f"参加工作时间为{t2}。"

        rdbd1 = f"    {name}同志档案中最早形成的《{a2}》中记载的出生年月为：{a1}，其档案中有记载{a3}的情况。{csny_alt}《{b2}》中记载的参加工作时间为：{b1}，其档案中有记载为{b3}的情况。{cjgz_alt}"

        if jl01_from_page:
            jl02 = f"    根据《中共中央组织部、人事部、公安部关于认真做好干部出生日期管理工作的通知》（组通字〔2006〕41号）及中共中央组织部《组工通讯》第10期相关规定，经集体研究，综合研判，建议{name}同志出生年月按{a1}进行认定。参加工作时间按{b1}进行认定。\r\n"
        else:
            jl02 = f"    根据《中共中央组织部、人事部、公安部关于认真做好干部出生日期管理工作的通知》（组通字〔2006〕41号）及中共中央组织部《组工通讯》第10期相关规定，建议{name}同志出生年月按{a1}进行认定。参加工作时间按{b1}进行认定。"

        rdbd2 = f"\n    经综合研判，{e5}，中共盘州市教育局党组会议研究，认定{name}同志出生年月为{a1}，参加工作时间为{b1}。"

        return JsonResponse({"code": 0, "RDBd1": rdbd1, "JL02": jl02, "RDBd2": rdbd2})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})


# ==================== 涂改记录 ====================


@login_required
@csrf_exempt
def alteration_save_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except:
        return JsonResponse({"code": 400})

    rsid = data.get("rsid")
    rows = data.get("rows", [])
    if not rsid or not rows:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        for row in rows:
            xh = row.get("序号", 0)
            lx = row.get("涂改类型", "")
            name = row.get("被涂改材料名称", "")
            ct = row.get("被涂改材料形成时间", "")
            tt = row.get("涂改后时间", "")
            cursor.execute(
                "SELECT COUNT(*) FROM YW_Alteration WHERE rsid=? AND xh=?",
                (int(rsid), int(xh)),
            )
            if cursor.fetchone()[0] > 0:
                cursor.execute(
                    "UPDATE YW_Alteration SET lx=?, name=?, createTime=?, totalTime=? WHERE rsid=? AND xh=?",
                    (lx, name, ct, tt, int(rsid), int(xh)),
                )
            else:
                cursor.execute(
                    "INSERT INTO YW_Alteration (rsid, xh, lx, name, createTime, totalTime) VALUES (?,?,?,?,?,?)",
                    (int(rsid), int(xh), lx, name, ct, tt),
                )
        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "保存成功"})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
def alteration_delete_api(request):
    rsid = request.GET.get("rsid", "")
    xh = request.GET.get("xh", "")
    if not rsid or not xh:
        return JsonResponse({"code": 400})
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "DELETE FROM YW_Alteration WHERE rsid=? AND xh=?", (int(rsid), int(xh))
        )
        cursor.execute(
            "UPDATE YW_Alteration SET xh=xh-1 WHERE rsid=? AND xh>?",
            (int(rsid), int(xh)),
        )
        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "删除成功"})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
def alteration_clear_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM YW_Alteration WHERE rsid=?", (int(rsid),))
        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "清空成功"})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
def identify_export_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    template_path = os.path.join(
        settings.BASE_DIR, "static", "excel_templates", "认定表.xlsx"
    )
    if not os.path.exists(template_path):
        return JsonResponse({"code": 500, "msg": "模板不存在"})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT XM, JOBUNIT FROM RS_INFO WHERE RSID=?", (int(rsid),))
        person = cursor.fetchone()
        cursor.execute(
            "SELECT a.rdyj, a.zzbmyj FROM YW_ZXSHDJ a WHERE a.RSID = ?", (int(rsid),)
        )
        data_row = cursor.fetchone()
        cursor.close()

        if not person:
            return JsonResponse({"code": 404})

        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        merged = list(ws.merged_cells.ranges)
        for mr in merged:
            ws.unmerge_cells(str(mr))

        ws.cell(row=3, column=2).value = person[0] or ""
        ws.cell(row=3, column=5).value = person[1] or ""

        if data_row:
            rdyj = (data_row[0] or "").replace("卍", "")
            zzbmyj = data_row[1] or ""
            ws.cell(row=4, column=2).value = rdyj
            ws.cell(row=24, column=2).value = zzbmyj

        for mr in merged:
            ws.merge_cells(str(mr))

        fn = f"{person[0]}_{rsid}_认定表.xlsx"
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(
            buf,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        encoded_fn = quote(fn.encode("utf-8"))
        response["Content-Disposition"] = (
            f"attachment; filename=\"{encoded_fn}\"; filename*=UTF-8''{encoded_fn}"
        )
        return response
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()
