from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.conf import settings
import json
import logging
import os
from io import BytesIO
import xlrd
from xlutils.copy import copy
from urllib.parse import quote
import openpyxl
from main.db_utils import _get_conn


@login_required
def person_position_view(request):
    return render(request, "archives/person_position.html")


@login_required
def position_data_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT sXH AS 序号, RZSJ AS 任职时间, MZSJ AS 免职时间,
                   BMmc AS 部门, ZW AS 职务, rzwh AS 批准文电号
            FROM YW_ZWBD WHERE RSID=? ORDER BY sXH
        """,
            (int(rsid),),
        )
        cols = [col[0] for col in cursor.description]
        rows = [dict(zip(cols, r)) for r in cursor.fetchall()]
        cursor.close()
        return JsonResponse({"code": 0, "data": rows})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@csrf_exempt
def position_save_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except:
        return JsonResponse({"code": 400})

    rsid = data.get("rsid")
    rows = data.get("rows", [])

    if not rows:
        return JsonResponse({"code": 400, "msg": "无数据"})

    values = []
    for row in rows:
        sxh = row.get("序号", 0)
        rzsj = (row.get("任职时间") or "").replace("'", "''")
        mzsj = (row.get("免职时间") or "").replace("'", "''")
        bm = (row.get("部门") or "").replace("'", "''")
        zw = (row.get("职务") or "").replace("'", "''")
        rzwh = (row.get("批准文电号") or "").replace("'", "''")
        values.append(f"({rsid},{sxh},'{rzsj}','{mzsj}','{bm}','{zw}','{rzwh}')")

    sql = f"""
        MERGE YW_ZWBD AS t
        USING (VALUES {",".join(values)}) AS s(RSID, sXH, RZSJ, MZSJ, BMmc, ZW, rzwh)
        ON t.RSID = s.RSID AND t.sXH = s.sXH
        WHEN MATCHED THEN UPDATE SET RZSJ=s.RZSJ, MZSJ=s.MZSJ, BMmc=s.BMmc, ZW=s.ZW, rzwh=s.rzwh
        WHEN NOT MATCHED THEN INSERT (RSID, sXH, RZSJ, MZSJ, BMmc, ZW, rzwh) VALUES (s.RSID, s.sXH, s.RZSJ, s.MZSJ, s.BMmc, s.ZW, s.rzwh);
    """

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(sql)
        conn.commit()
        return JsonResponse({"code": 0, "msg": "保存成功"})
    except Exception as e:
        logger.error(f"职务保存失败: {e}")
        if conn:
            try:
                conn.rollback()
            except:
                pass
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass


@login_required
def position_export_api(request):
    rsid = request.GET.get("rsid", "")
    if not rsid:
        return JsonResponse({"code": 400})

    template_path = os.path.join(
        settings.BASE_DIR, "static", "excel_templates", "9_2_1.xlsx"
    )

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT XM, XB, CSNY, JG FROM RS_INFO WHERE RSID=?", (int(rsid),)
        )
        person = cursor.fetchone()
        cursor.execute(
            "SELECT RZSJ, MZSJ, BMmc, ZW, RZWH FROM YW_ZWBD WHERE RSID=? ORDER BY SXH",
            (int(rsid),),
        )
        rows = cursor.fetchall()
        cursor.close()

        if not person:
            return JsonResponse({"code": 404})

        wb = openpyxl.load_workbook(template_path)
        ws = wb.active

        merged = list(ws.merged_cells.ranges)
        for mr in merged:
            ws.unmerge_cells(str(mr))

        ws.cell(row=2, column=2).value = person[0] or ""
        ws.cell(row=2, column=5).value = person[1] or ""
        ws.cell(row=2, column=7).value = person[2] or ""
        ws.cell(row=2, column=11).value = person[3] or ""

        for i, row in enumerate(rows):
            r = i + 4
            ws.cell(row=r, column=1).value = row[0] or ""
            ws.cell(row=r, column=2).value = row[1] or ""
            ws.cell(row=r, column=3).value = row[2] or ""
            ws.cell(row=r, column=7).value = row[3] or ""
            ws.cell(row=r, column=10).value = row[4] or ""

        for mr in merged:
            ws.merge_cells(str(mr))

        export_dir = os.path.join(settings.MEDIA_ROOT, "exports")
        os.makedirs(export_dir, exist_ok=True)
        fn = f"{person[0]}_{rsid}_9-2-1.xlsx"
        filepath = os.path.join(export_dir, fn)
        wb.save(filepath)

        return JsonResponse({"code": 0, "url": f"/media/exports/{fn}"})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass
