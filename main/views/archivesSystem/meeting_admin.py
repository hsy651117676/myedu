"""
meeting_admin.py
党组会管理（管理员）
"""

import hashlib
import json
import logging
import os
import openpyxl
from datetime import datetime
from urllib.parse import quote
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.core.cache import cache
from django.http import FileResponse, Http404, JsonResponse
from django.shortcuts import render
from django.views.decorators.csrf import csrf_exempt
from django.http import HttpResponse

from main.utils import _get_conn
from main.utils.decorators import admin_required

logger = logging.getLogger(__name__)


@login_required
@admin_required
def page(request):
    return render(request, "archivesSystem/meeting_admin.html")


@login_required
@admin_required
def batch_list_api(request):
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT batch_id, batch_name, batch_date, leader, editable FROM z_meeting_batch ORDER BY batch_date DESC"
        )
        data = [
            {
                "batch_id": r[0],
                "batch_name": r[1],
                "batch_date": r[2],
                "leader": r[3],
                "editable": r[4],
            }
            for r in cursor.fetchall()
        ]
        cursor.close()
        return JsonResponse({"code": 0, "data": data})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
def all_batches_api(request):
    """所有届次"""
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM z_meeting_batch ORDER BY batch_id DESC")
        cols = [c[0] for c in cursor.description]
        rows = [dict(zip(cols, r)) for r in cursor.fetchall()]
        cursor.close()
        return JsonResponse({"code": 0, "data": rows})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
def list_api(request):
    """人员列表"""
    batch_id = request.GET.get("batch_id", "")
    name = request.GET.get("name", "")
    page = int(request.GET.get("page", 1))
    page_size = int(request.GET.get("pageSize", 30))

    cache_key = f"meeting_admin:list:{hashlib.md5(f'{batch_id}{name}{page}{page_size}'.encode()).hexdigest()}"
    cached = cache.get(cache_key)
    if cached is not None:
        return JsonResponse(
            {"code": 0, "data": cached.get("data", []), "total": cached.get("total", 0)}
        )

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        where = "WHERE 1=1"
        params = []

        if batch_id == "null":
            where += " AND z.batch_id IS NULL"
        elif batch_id:
            where += " AND z.batch_id = ?"
            params.append(int(batch_id))
        if name:
            where += " AND z.rsid IN (SELECT RSID FROM RS_INFO WHERE XM LIKE ?)"
            params.append(f"%{name}%")

        cursor.execute(f"SELECT COUNT(*) FROM z_meeting z {where}", params)
        total = cursor.fetchone()[0]

        offset = (page - 1) * page_size
        cursor.execute(
            f"""
            SELECT z.*, 
                   (SELECT XM FROM RS_INFO WHERE RSID = z.rsid) AS XM,
                   (SELECT batch_name FROM z_meeting_batch WHERE batch_id = z.batch_id) AS batch_name,
                   (SELECT editable FROM z_meeting_batch WHERE batch_id = z.batch_id) AS editable
            FROM z_meeting z
            {where}
            ORDER BY z.rsid
            OFFSET {offset} ROWS FETCH NEXT {page_size} ROWS ONLY
        """,
            params,
        )
        cols = [c[0] for c in cursor.description]
        rows = [dict(zip(cols, r)) for r in cursor.fetchall()]
        cursor.close()

        result = {"data": rows, "total": total}
        cache.set(cache_key, result, 1800)
        return JsonResponse({"code": 0, "data": rows, "total": total})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
@csrf_exempt
def save_batch_api(request):
    """新增/更新届次"""
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({"code": 400})

    batch_id = data.get("batch_id")
    batch_name = data.get("batch_name", "")
    batch_date = data.get("batch_date", "")
    leader = data.get("leader", "")
    applicant = data.get("applicant", "")
    editable = 1 if data.get("editable", 1) else 0
    batch_status = data.get("batch_status", "进行中")

    if not batch_name:
        return JsonResponse({"code": 400, "msg": "届次名称不能为空"})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        if batch_id:
            cursor.execute(
                "UPDATE z_meeting_batch SET batch_name=?, batch_date=?, leader=?, applicant=?, editable=?, batch_status=? WHERE batch_id=?",
                (
                    batch_name,
                    batch_date,
                    leader,
                    applicant,
                    editable,
                    batch_status,
                    int(batch_id),
                ),
            )
        else:
            cursor.execute(
                "INSERT INTO z_meeting_batch (batch_name, batch_date, leader, applicant, editable, batch_status) VALUES (?,?,?,?,?,?)",
                (batch_name, batch_date, leader, applicant, editable, batch_status),
            )
        conn.commit()
        cursor.close()
        cache.delete_pattern("meeting_admin:list:*")
        return JsonResponse({"code": 0, "msg": "保存成功"})
    except Exception as e:
        if conn:
            conn.rollback()
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
@csrf_exempt
def delete_batch_api(request):
    """删除届次"""
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({"code": 400})
    batch_id = data.get("batch_id")
    if not batch_id:
        return JsonResponse({"code": 400})
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT COUNT(*) FROM z_meeting WHERE batch_id=?", (int(batch_id),)
        )
        if cursor.fetchone()[0] > 0:
            cursor.close()
            return JsonResponse({"code": 400, "msg": "该届次下有人员，无法删除"})
        cursor.execute("DELETE FROM z_meeting_batch WHERE batch_id=?", (int(batch_id),))
        conn.commit()
        cursor.close()
        cache.delete_pattern("meeting_admin:list:*")
        return JsonResponse({"code": 0, "msg": "删除成功"})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
@csrf_exempt
def upload_file_api(request):
    """上传届次文件"""
    if request.method != "POST":
        return JsonResponse({"code": 405})
    batch_id = request.POST.get("batch_id")
    file_type = request.POST.get("file_type")
    uploaded_file = request.FILES.get("file")
    if not batch_id or not file_type or not uploaded_file:
        return JsonResponse({"code": 400, "msg": "缺少参数"})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT file_path FROM z_meeting_batch WHERE batch_id=?", (int(batch_id),)
        )
        row = cursor.fetchone()
        base_path = row[0] if row and row[0] else "/mnt/data/meeting_files"
        os.makedirs(base_path, exist_ok=True)

        ext = uploaded_file.name.rsplit(".", 1)[-1] if "." in uploaded_file.name else ""
        save_name = f"{batch_id}_{file_type}.{ext}"
        full_path = os.path.join(base_path, save_name)
        with open(full_path, "wb+") as f:
            for chunk in uploaded_file.chunks():
                f.write(chunk)

        col = f"file_name_{file_type}"
        cursor.execute(
            f"UPDATE z_meeting_batch SET file_path=?, {col}=? WHERE batch_id=?",
            (base_path, uploaded_file.name, int(batch_id)),
        )
        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "上传成功"})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
def download_file_api(request):
    """下载届次文件"""
    batch_id = request.GET.get("id")
    file_type = request.GET.get("type")
    if not batch_id or not file_type:
        raise Http404

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        col = f"file_name_{file_type}"
        cursor.execute(
            f"SELECT file_path, {col} FROM z_meeting_batch WHERE batch_id=?",
            (int(batch_id),),
        )
        row = cursor.fetchone()
        cursor.close()
        if not row or not row[0] or not row[1]:
            raise Http404
        ext = row[1].rsplit(".", 1)[-1] if "." in row[1] else ""
        full = os.path.join(row[0], f"{batch_id}_{file_type}.{ext}")
        if os.path.exists(full):
            return FileResponse(open(full, "rb"), filename=row[1])
        raise Http404
    except Http404:
        raise
    except Exception:
        raise Http404
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
@csrf_exempt
def batch_file_api(request):
    """批量上会申请"""
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({"code": 400})
    rsids = data.get("rsids", [])
    batch_id = data.get("batch_id", "")
    if not rsids or not batch_id:
        return JsonResponse({"code": 400})
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        for rsid in rsids:
            cursor.execute(
                "UPDATE z_meeting SET batch_id=? WHERE rsid=?",
                (int(batch_id), int(rsid)),
            )
        conn.commit()
        cursor.close()
        cache.delete_pattern("meeting:list:*")
        cache.delete_pattern("meeting_admin:list:*")
        return JsonResponse({"code": 0, "msg": f"成功上会 {len(rsids)} 人"})
    except Exception as e:
        if conn:
            conn.rollback()
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
@csrf_exempt
def batch_delete_api(request):
    """批量删除"""
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({"code": 400})
    rsids = data.get("rsids", [])
    if not rsids:
        return JsonResponse({"code": 400})
    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        for rsid in rsids:
            cursor.execute("DELETE FROM z_meeting WHERE rsid=?", (int(rsid),))
        conn.commit()
        cursor.close()
        cache.delete_pattern("meeting:list:*")
        cache.delete_pattern("meeting_admin:list:*")
        return JsonResponse({"code": 0, "msg": f"成功删除 {len(rsids)} 人"})
    except Exception as e:
        if conn:
            conn.rollback()
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
@csrf_exempt
def admin_save_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({"code": 400})
    rsid = data.get("rsid")
    if not rsid:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        fields = [
            "school_name",
            "identification_type",
            "birth_summary",
            "birth_material_name",
            "birth_material_time",
            "birth_material_date",
            "birth_alteration",
            "birth_decision_date",
            "work_summary",
            "work_material_name",
            "work_material_time",
            "work_material_date",
            "work_alteration",
            "work_decision_date",
            "party_join_time",
            "party_material_date",
            "party_decision_date",
            "meeting_type",
            "remark",
        ]
        values = [data.get(f, "") for f in fields]

        cursor.execute("SELECT rsid FROM z_meeting WHERE rsid=?", (int(rsid),))
        if cursor.fetchone():
            set_clause = ",".join([f"{f}=?" for f in fields])
            cursor.execute(
                f"UPDATE z_meeting SET {set_clause} WHERE rsid=?", values + [int(rsid)]
            )
        else:
            placeholders = ",".join(["?"] * len(fields))
            cursor.execute(
                f"INSERT INTO z_meeting (rsid,{','.join(fields)}) VALUES (?,{placeholders})",
                [int(rsid)] + values,
            )
        conn.commit()
        cursor.close()
        cache.delete_pattern("meeting:list:*")
        cache.delete_pattern("meeting_admin:list:*")
        return JsonResponse({"code": 0, "msg": "保存成功"})
    except Exception as e:
        if conn:
            conn.rollback()
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@admin_required
def export_excel_api(request):
    batch_id = request.GET.get("batch_id", "")

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        where = "WHERE 1=1"
        params = []
        if batch_id == "null":
            where += " AND z.batch_id IS NULL"
        elif batch_id:
            where += " AND z.batch_id = ?"
            params.append(int(batch_id))

        cursor.execute(
            f"""
            SELECT z.*, 
                   (SELECT XM FROM RS_INFO WHERE RSID = z.rsid) AS XM,
                   (SELECT batch_name FROM z_meeting_batch WHERE batch_id = z.batch_id) AS batch_name
            FROM z_meeting z {where}
            ORDER BY z.rsid
        """,
            params,
        )
        cols = [c[0] for c in cursor.description]
        rows = [dict(zip(cols, r)) for r in cursor.fetchall()]
        cursor.close()
    finally:
        if conn:
            conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "党组会名单"

    headers = [
        "姓名",
        "学校名称",
        "认定类型",
        "出生记载",
        "最早材料名称",
        "记载出生时间",
        "材料形成时间",
        "认定时间",
        "涂改情况",
        "参工记载",
        "依据材料名称",
        "记载参工时间",
        "材料形成时间",
        "认定时间",
        "涂改情况",
        "入党时间",
        "材料时间",
        "认定时间",
        "上会类型",
        "届次",
        "备注",
    ]

    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)

    for i, r in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=r.get("XM", ""))
        ws.cell(row=i, column=2, value=r.get("school_name", ""))
        ws.cell(row=i, column=3, value=r.get("identification_type", ""))
        ws.cell(row=i, column=4, value=r.get("birth_summary", ""))
        ws.cell(row=i, column=5, value=r.get("birth_material_name", ""))
        ws.cell(row=i, column=6, value=r.get("birth_material_time", ""))
        ws.cell(row=i, column=7, value=r.get("birth_material_date", ""))
        ws.cell(row=i, column=8, value=r.get("birth_decision_date", ""))
        ws.cell(row=i, column=9, value=r.get("birth_alteration", ""))
        ws.cell(row=i, column=10, value=r.get("work_summary", ""))
        ws.cell(row=i, column=11, value=r.get("work_material_name", ""))
        ws.cell(row=i, column=12, value=r.get("work_material_time", ""))
        ws.cell(row=i, column=13, value=r.get("work_material_date", ""))
        ws.cell(row=i, column=14, value=r.get("work_decision_date", ""))
        ws.cell(row=i, column=15, value=r.get("work_alteration", ""))
        ws.cell(row=i, column=16, value=r.get("party_join_time", ""))
        ws.cell(row=i, column=17, value=r.get("party_material_date", ""))
        ws.cell(row=i, column=18, value=r.get("party_decision_date", ""))
        ws.cell(row=i, column=19, value=r.get("meeting_type", ""))
        ws.cell(row=i, column=20, value=r.get("batch_name", ""))
        ws.cell(row=i, column=21, value=r.get("remark", ""))

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"党组会名单_{datetime.now().strftime('%Y%m%d')}.xlsx"
    resp = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return resp


@login_required
@admin_required
def export_template_api(request):
    return JsonResponse({"code": 0, "msg": "模板功能开发中"})
