"""
学生档案管理
"""

import json
import os
import logging
from urllib.parse import quote
from io import BytesIO
from django.conf import settings
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from main.db_utils import _get_conn

logger = logging.getLogger(__name__)

FIELDS = [
    "School",
    "Name",
    "GraduationYear",
    "SerialNo",
    "TransferNo",
    "TransferOutDate",
    "TransferOutPlace",
    "TransferOutPerson",
    "Handler",
    "Remark",
]

EXCEL_TEMPLATE_PATH = os.path.join(
    settings.BASE_DIR, "static", "excel_templates", "干部档案传递单.xlsx"
)


@login_required
def page(request):
    return render(request, "business/student_archives.html")


@login_required
def list_api(request):
    keyword = request.GET.get("keyword", "").strip()
    page = int(request.GET.get("page", 1))
    page_size = int(request.GET.get("pageSize", 30))

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        where = "WHERE IsActive=1"
        params = []
        if keyword:
            where += " AND (Name LIKE ? OR School LIKE ?)"
            kw = f"%{keyword}%"
            params.extend([kw, kw])

        cursor.execute(f"SELECT COUNT(*) FROM StudentArchives {where}", params)
        total = cursor.fetchone()[0]

        offset = (page - 1) * page_size
        cursor.execute(
            f"""
            SELECT * FROM StudentArchives {where}
            ORDER BY ID DESC
            OFFSET {offset} ROWS FETCH NEXT {page_size} ROWS ONLY
        """,
            params,
        )
        cols = [c[0] for c in cursor.description]
        rows = [dict(zip(cols, r)) for r in cursor.fetchall()]
        cursor.close()
        return JsonResponse({"code": 0, "data": rows, "total": total})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@csrf_exempt
def save_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except:
        return JsonResponse({"code": 400})

    record_id = data.get("ID")
    if not data.get("Name"):
        return JsonResponse({"code": 400, "msg": "姓名不能为空"})

    values = [data.get(f, "") for f in FIELDS]

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        if record_id:
            set_clause = ",".join([f"{f}=?" for f in FIELDS])
            cursor.execute(
                f"UPDATE StudentArchives SET {set_clause}, UpdateTime=GETDATE() WHERE ID=?",
                values + [int(record_id)],
            )
        else:
            placeholders = ",".join(["?"] * len(FIELDS))
            cursor.execute(
                f"INSERT INTO StudentArchives ({','.join(FIELDS)}) VALUES ({placeholders})",
                values,
            )
        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "保存成功"})
    except Exception as e:
        if conn:
            conn.rollback()
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@csrf_exempt
def delete_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
    except:
        return JsonResponse({"code": 400})
    record_id = data.get("ID")
    if not record_id:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE StudentArchives SET IsActive=0, UpdateTime=GETDATE() WHERE ID=?",
            (int(record_id),),
        )
        conn.commit()
        cursor.close()
        return JsonResponse({"code": 0, "msg": "删除成功"})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
def print_view(request):
    """打印预览页面"""
    record_id = request.GET.get("id", "")
    if not record_id:
        return HttpResponse("缺少ID", status=400)

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT * FROM StudentArchives WHERE ID=? AND IsActive=1", (int(record_id),)
        )
        cols = [c[0] for c in cursor.description]
        row = cursor.fetchone()
        cursor.close()
        if not row:
            return HttpResponse("记录不存在", status=404)
        data = dict(zip(cols, row))
    finally:
        if conn:
            conn.close()

    context = {
        "zdwh": data.get("TransferNo", ""),
        "fb": data.get("Name", ""),
        "person_job": data.get("School", ""),
        "zwdw": data.get("TransferOutPlace", ""),
        "zdsj_display": data.get("TransferOutDate", ""),
        "rid": record_id,
    }
    return render(request, "business/student_transfer_print.html", context)


@login_required
def print_excel_api(request):
    """导出Excel"""
    record_id = request.GET.get("id", "")
    if not record_id:
        return HttpResponse("缺少ID", status=400)

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT * FROM StudentArchives WHERE ID=? AND IsActive=1", (int(record_id),)
        )
        cols = [c[0] for c in cursor.description]
        row = cursor.fetchone()
        cursor.close()
        if not row:
            return HttpResponse("记录不存在", status=404)
        data = dict(zip(cols, row))
    finally:
        if conn:
            conn.close()

    import openpyxl

    wb = openpyxl.load_workbook(EXCEL_TEMPLATE_PATH)
    sheet = wb["档案转递单"]

    zdwh = data.get("TransferNo", "")
    fb = data.get("Name", "")
    person_job = data.get("School", "")
    zwdw = data.get("TransferOutPlace", "")
    zdsj_display = data.get("TransferOutDate", "")

    sheet.cell(row=2, column=3).value = zdwh
    sheet.cell(row=7, column=3).value = zdwh
    sheet.cell(row=4, column=1).value = fb
    sheet.cell(row=4, column=2).value = person_job
    sheet.cell(row=4, column=3).value = zwdw
    sheet.cell(row=8, column=1).value = f"{zwdw}："
    sheet.cell(
        row=9, column=1
    ).value = f"  {fb}等同志的档案材料转出，请按档案内所列目录清点查收，并将回执退回。"
    sheet.cell(row=11, column=3).value = zdsj_display
    sheet.cell(row=13, column=1).value = fb
    sheet.cell(row=13, column=2).value = person_job
    sheet.cell(row=13, column=3).value = zwdw

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"学生档案转递单_{fb}.xlsx"
    resp = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return resp


@login_required
def print_pdf_api(request):
    """导出PDF"""
    record_id = request.GET.get("id", "")
    if not record_id:
        return HttpResponse("缺少ID", status=400)

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT * FROM StudentArchives WHERE ID=? AND IsActive=1", (int(record_id),)
        )
        cols = [c[0] for c in cursor.description]
        row = cursor.fetchone()
        cursor.close()
        if not row:
            return HttpResponse("记录不存在", status=404)
        data = dict(zip(cols, row))
    finally:
        if conn:
            conn.close()

    from weasyprint import HTML

    zdwh = data.get("TransferNo", "")
    fb = data.get("Name", "")
    person_job = data.get("School", "")
    zwdw = data.get("TransferOutPlace", "")
    zdsj_display = data.get("TransferOutDate", "")

    html = f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><style>
@page{{size:A4 portrait;margin:15mm 12mm 15mm 12mm;}}
body{{font-family:"SimSun","宋体",sans-serif;font-size:14pt;margin:0;}}
h1{{text-align:center;font-size:20pt;font-weight:bold;margin:0 0 6px 0;}}
h2{{text-align:center;font-size:20pt;font-weight:bold;margin:10px 0 6px 0;}}
.file-no{{text-align:center;font-size:14pt;margin:0 0 6px 0;}}
table{{border-collapse:collapse;width:100%;}}
td{{border:1px solid #000;padding:6px 8px;font-size:14pt;vertical-align:middle;text-align:center;}}
.sep{{text-align:center;font-size:14pt;margin:6px 0;letter-spacing:4px;}}
.content{{border:1px solid #000;padding:8px;margin:6px 0;font-size:14pt;line-height:1.8;text-align:left;}}
.right{{text-align:right;font-size:14pt;margin:4px 0;}}
.return-title{{text-align:center;font-weight:bold;font-size:16pt;margin:8px 0;}}
.footer{{font-size:12pt;margin-top:12px;line-height:1.6;}}
.c1{{width:15%;}}.c2{{width:42%;}}.c3{{width:28%;}}.c4{{width:15%;}}
</style></head><body>

<h1>干部档案传递存根</h1>
<div class="file-no">{zdwh}</div>
<table>
<tr><td class="c1">姓名</td><td class="c2">工作单位</td><td class="c3">调往何单位</td><td class="c4">档案卷数</td></tr>
<tr><td class="c1">{fb}</td><td class="c2">{person_job}</td><td class="c3">{zwdw}</td><td class="c4">壹卷</td></tr>
</table>

<div class="sep">…………………………………………………………………………………………</div>
<h2>干部档案传递通知单</h2>
<div class="file-no">{zdwh}</div>
<div style="font-size:14pt;">{zwdw}：</div>
<div class="content">
  {fb}等同志的档案材料转出，请按档案内所列目录清点查收，并将回执退回。
</div>
<div class="right">盘州市教育局档案室</div>
<div class="right">{zdsj_display}</div>


<table>
<tr><td class="c1">姓名</td><td class="c2">工作单位</td><td class="c3">调往何单位</td><td class="c4">档案卷数</td></tr>
<tr><td class="c1">{fb}</td><td class="c2">{person_job}</td><td class="c3">{zwdw}</td><td class="c4">壹卷</td></tr>
</table>

<div style="font-size:14pt;margin:4px 0;">盘州市教育局：</div>
<div class="return-title">回&emsp;执</div>
<div class="content" style="font-size:12pt;line-height:2;">
&emsp;&emsp;你处&emsp;&emsp;年&emsp;月&emsp;日转来的第&emsp;&emsp;号干部档案转递通知单中所列的&emsp;&emsp;&emsp;&emsp;&emsp;&emsp;等&emsp;&emsp;名同志档案共&emsp;&emsp;卷，我处已于&emsp;&emsp;年&emsp;月&emsp;日收到，经清点无误，现将回执退回，请查收。
</div>
<div style="margin:8px 0;">
  <span>收件人签字：</span>
  <span style="margin-left:40px;">收件单位（盖章）</span>
  <span style="margin-left:60px;">年&emsp;月&emsp;日</span>
</div>
<div class="footer">
  回执邮寄地址：贵州省六盘水市盘州市红果街道盘西新城纵五路教育局大楼517档案室<br>
  联系电话：0858-3634122
</div>
</body></html>"""

    buf = BytesIO()
    HTML(string=html).write_pdf(buf)
    buf.seek(0)
    filename = f"学生档案转递单_{fb}.pdf"
    resp = HttpResponse(buf.read(), content_type="application/pdf")
    resp["Content-Disposition"] = f'inline; filename="{quote(filename)}"'
    return resp
