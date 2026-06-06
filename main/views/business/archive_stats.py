"""
档案库存统计
"""

import json
from io import BytesIO
from urllib.parse import quote
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.core.cache import cache
from main.db_utils import _get_conn


def _get_unit_stats(cursor, tids):
    """根据tid列表统计单位人数，自动展开分组ID为子单位ID"""
    if not tids:
        return []
    placeholders = ",".join(["?"] * len(tids))
    cursor.execute(
        f"SELECT TID FROM BMGL WHERE TID IN ({placeholders}) AND PID = -1", tids
    )
    group_tids = set(str(r[0]) for r in cursor.fetchall())

    expanded = []
    for tid in tids:
        if str(tid) in group_tids:
            cursor.execute("SELECT TID FROM BMGL WHERE PID = ? AND PID > 0", (tid,))
            expanded.extend(str(r[0]) for r in cursor.fetchall())
        else:
            expanded.append(str(tid))

    if not expanded:
        return []

    placeholders = ",".join(["?"] * len(expanded))
    cursor.execute(
        f"""
        SELECT c.TID, c.TNAME, COUNT(DISTINCT r.RSID) AS CNT
        FROM RS_INFO r
        INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
        INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
        WHERE c.TID IN ({placeholders}) AND c.PID > 0
        GROUP BY c.TID, c.TNAME
        ORDER BY CNT DESC
    """,
        expanded,
    )
    return [{"tid": r[0], "tname": r[1], "人数": r[2]} for r in cursor.fetchall()]


def _get_persons_by_tid(cursor, tid):
    """根据tid获取人员列表，分组ID自动展开"""
    cursor.execute("SELECT PID FROM BMGL WHERE TID = ?", (tid,))
    row = cursor.fetchone()
    if row and row[0] == -1:
        cursor.execute(
            """
            SELECT r.RSID, r.XM, r.XB, r.CSNY, r.ZZMM, r.JOBUNIT, r.RYBH
            FROM RS_INFO r
            INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
            INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
            WHERE c.PID = ? AND c.PID > 0
            ORDER BY r.RSID
        """,
            (tid,),
        )
    else:
        cursor.execute(
            """
            SELECT r.RSID, r.XM, r.XB, r.CSNY, r.ZZMM, r.JOBUNIT, r.RYBH
            FROM RS_INFO r
            INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
            INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
            WHERE c.TID = ? AND c.PID > 0
            ORDER BY r.RSID
        """,
            (tid,),
        )
    cols = [c[0] for c in cursor.description]
    return [dict(zip(cols, r)) for r in cursor.fetchall()]


def _get_summary_stats(cursor, unit_tids):
    """根据单位tid列表获取汇总统计"""
    if not unit_tids:
        return {"总人数": 0, "男": 0, "女": 0, "性别未知": 0}, [], []

    placeholders = ",".join(["?"] * len(unit_tids))
    tids = [str(t) for t in unit_tids]

    # 汇总（含性别未知）
    cursor.execute(
        f"""
        SELECT COUNT(DISTINCT r.RSID),
               SUM(CASE WHEN r.XB = '男' THEN 1 ELSE 0 END),
               SUM(CASE WHEN r.XB = '女' THEN 1 ELSE 0 END),
               SUM(CASE WHEN r.XB IS NULL OR r.XB = '' OR r.XB NOT IN ('男','女') THEN 1 ELSE 0 END)
        FROM RS_INFO r
        INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
        INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
        WHERE c.TID IN ({placeholders}) AND c.PID > 0
    """,
        tids,
    )
    row = cursor.fetchone()
    summary = {
        "总人数": row[0] or 0,
        "男": row[1] or 0,
        "女": row[2] or 0,
        "性别未知": row[3] or 0,
    }

    # 年龄分布
    age_ranges = [
        ("35岁以下", 0, 35),
        ("36-40岁", 36, 40),
        ("41-45岁", 41, 45),
        ("46-50岁", 46, 50),
        ("51-55岁", 51, 55),
        ("56-60岁", 56, 60),
        ("60岁以上", 61, 200),
    ]
    age_data = []
    for name, lo, hi in age_ranges:
        cursor.execute(
            f"""
            SELECT COUNT(DISTINCT r.RSID),
                   SUM(CASE WHEN r.XB='男' THEN 1 ELSE 0 END),
                   SUM(CASE WHEN r.XB='女' THEN 1 ELSE 0 END)
            FROM RS_INFO r
            INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
            INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
            WHERE c.TID IN ({placeholders}) AND c.PID > 0
              AND TRY_CAST(LEFT(r.CSNY,4) AS INT) BETWEEN YEAR(GETDATE())-{hi} AND YEAR(GETDATE())-{lo}
        """,
            tids,
        )
        row = cursor.fetchone()
        age_data.append(
            {"name": name, "人数": row[0] or 0, "男": row[1] or 0, "女": row[2] or 0}
        )

    # 年龄未知
    cursor.execute(
        f"""
        SELECT COUNT(DISTINCT r.RSID),
               SUM(CASE WHEN r.XB='男' THEN 1 ELSE 0 END),
               SUM(CASE WHEN r.XB='女' THEN 1 ELSE 0 END)
        FROM RS_INFO r
        INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
        INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
        WHERE c.TID IN ({placeholders}) AND c.PID > 0
          AND (r.CSNY IS NULL OR r.CSNY = '' OR TRY_CAST(LEFT(r.CSNY,4) AS INT) IS NULL)
    """,
        tids,
    )
    row = cursor.fetchone()
    if row[0]:
        age_data.append(
            {"name": "未知", "人数": row[0] or 0, "男": row[1] or 0, "女": row[2] or 0}
        )

    # 工龄分布
    work_ranges = [
        ("5年以下", 0, 5),
        ("6-10年", 6, 10),
        ("11-15年", 11, 15),
        ("16-20年", 16, 20),
        ("21-25年", 21, 25),
        ("26-30年", 26, 30),
        ("30年以上", 31, 60),
    ]
    work_data = []
    for name, lo, hi in work_ranges:
        cursor.execute(
            f"""
            SELECT COUNT(DISTINCT r.RSID),
                   SUM(CASE WHEN r.XB='男' THEN 1 ELSE 0 END),
                   SUM(CASE WHEN r.XB='女' THEN 1 ELSE 0 END)
            FROM RS_INFO r
            INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
            INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
            WHERE c.TID IN ({placeholders}) AND c.PID > 0
              AND TRY_CAST(LEFT(r.WORKTIME,4) AS INT) BETWEEN YEAR(GETDATE())-{hi} AND YEAR(GETDATE())-{lo}
        """,
            tids,
        )
        row = cursor.fetchone()
        work_data.append(
            {"name": name, "人数": row[0] or 0, "男": row[1] or 0, "女": row[2] or 0}
        )

    # 工龄未知
    cursor.execute(
        f"""
        SELECT COUNT(DISTINCT r.RSID),
               SUM(CASE WHEN r.XB='男' THEN 1 ELSE 0 END),
               SUM(CASE WHEN r.XB='女' THEN 1 ELSE 0 END)
        FROM RS_INFO r
        INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
        INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
        WHERE c.TID IN ({placeholders}) AND c.PID > 0
          AND (r.WORKTIME IS NULL OR r.WORKTIME = '' OR TRY_CAST(LEFT(r.WORKTIME,4) AS INT) IS NULL)
    """,
        tids,
    )
    row = cursor.fetchone()
    if row[0]:
        work_data.append(
            {"name": "未知", "人数": row[0] or 0, "男": row[1] or 0, "女": row[2] or 0}
        )

    return summary, age_data, work_data


@login_required
def page(request):
    return render(request, "business/archive_stats.html")


@login_required
def stats_api(request):
    tid = request.GET.get("tid", "")
    cache_key = f"archive_stats:{tid}"
    cached = cache.get(cache_key)
    if cached is not None:
        return JsonResponse(cached)

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()

        if tid == "":
            cursor.execute("""
                SELECT p.TID, p.TNAME, COUNT(DISTINCT r.RSID) AS CNT
                FROM RS_INFO r
                INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
                INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
                INNER JOIN BMGL p ON c.PID = p.TID
                WHERE p.PID = -1
                GROUP BY p.TID, p.TNAME
                ORDER BY CNT DESC
            """)
            rows = [
                {"tid": r[0], "tname": r[1], "count": r[2]} for r in cursor.fetchall()
            ]
            total = sum(r["count"] for r in rows)
            cursor.close()
            result = JsonResponse({"code": 0, "groups": rows, "total": total})
            cache.set(cache_key, {"code": 0, "groups": rows, "total": total}, 3600)
            return result

        cursor.execute("SELECT PID FROM BMGL WHERE TID = ?", (tid,))
        row = cursor.fetchone()

        if row and row[0] == -1:
            cursor.execute(
                """
                SELECT c.TID, c.TNAME, COUNT(DISTINCT r.RSID) AS CNT
                FROM RS_INFO r
                INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
                INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
                WHERE c.PID = ? AND c.PID > 0
                GROUP BY c.TID, c.TNAME
                ORDER BY CNT DESC
            """,
                (tid,),
            )
            rows = [
                {"tid": r[0], "tname": r[1], "count": r[2]} for r in cursor.fetchall()
            ]
            cursor.close()
            result = JsonResponse({"code": 0, "children": rows})
            cache.set(cache_key, {"code": 0, "children": rows}, 3600)
            return result
        else:
            cursor.execute(
                """
                SELECT c.TID, c.TNAME, COUNT(DISTINCT r.RSID) AS CNT
                FROM RS_INFO r
                INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
                INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
                WHERE c.TID = ? AND c.PID > 0
                GROUP BY c.TID, c.TNAME
            """,
                (tid,),
            )
            r = cursor.fetchone()
            row_data = [{"tid": r[0], "tname": r[1], "count": r[2]}] if r else []
            cursor.close()
            result = JsonResponse({"code": 0, "children": row_data})
            cache.set(cache_key, {"code": 0, "children": row_data}, 3600)
            return result

    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
@csrf_exempt
def checked_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
        tids = data.get("tids", [])
    except:
        return JsonResponse({"code": 400})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        rows = _get_unit_stats(cursor, tids)
        if not rows:
            cursor.close()
            return JsonResponse(
                {
                    "code": 0,
                    "summary": {"总人数": 0, "男": 0, "女": 0, "性别未知": 0},
                    "age": [],
                    "work": [],
                }
            )

        unit_tids = [r["tid"] for r in rows]
        summary, age_data, work_data = _get_summary_stats(cursor, unit_tids)
        cursor.close()
        return JsonResponse(
            {"code": 0, "summary": summary, "age": age_data, "work": work_data}
        )
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
def detail_api(request):
    tid = request.GET.get("tid", "")
    if not tid:
        return JsonResponse({"code": 400})

    cache_key = f"archive_stats_detail:{tid}"
    cached = cache.get(cache_key)
    if cached is not None:
        return JsonResponse({"code": 0, "data": cached})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        rows = _get_persons_by_tid(cursor, tid)
        cursor.close()
        cache.set(cache_key, rows, 3600)
        return JsonResponse({"code": 0, "data": rows})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()


@login_required
def export_api(request):
    tids_str = request.GET.get("tids", "")
    tids = [t for t in tids_str.split(",") if t]

    cache_key = f"archive_stats_export:{tids_str}"
    cached = cache.get(cache_key)
    if cached is not None:
        return cached

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        if tids:
            rows = _get_unit_stats(cursor, tids)
        else:
            cursor.execute("""
                SELECT p.TNAME, COUNT(DISTINCT r.RSID) AS CNT
                FROM RS_INFO r
                INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
                INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
                INNER JOIN BMGL p ON c.PID = p.TID
                WHERE p.PID = -1
                GROUP BY p.TNAME
                ORDER BY CNT DESC
            """)
            rows = [{"tname": r[0], "人数": r[1]} for r in cursor.fetchall()]
        cursor.close()
    finally:
        if conn:
            conn.close()

    import openpyxl
    from datetime import datetime

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "档案库存统计"
    ws.cell(row=1, column=1, value="单位")
    ws.cell(row=1, column=2, value="人数")
    for i, r in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=r.get("tname", ""))
        ws.cell(row=i, column=2, value=r.get("人数", 0))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"档案库存统计_{datetime.now().strftime('%Y%m%d')}.xlsx"
    resp = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    cache.set(cache_key, resp, 3600)
    return resp


@login_required
def export_detail_api(request):
    tids_str = request.GET.get("tids", "")
    tids = [t for t in tids_str.split(",") if t]

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        rows = _get_unit_stats(cursor, tids)
        if not rows:
            return HttpResponse("无数据", status=404)
        unit_tids = [str(r["tid"]) for r in rows]
        placeholders = ",".join(["?"] * len(unit_tids))
        cursor.execute(
            f"""
            SELECT r.XM AS 姓名, r.XB AS 性别, r.CSNY AS 出生年月, r.ZZMM AS 政治面貌,
                r.JOBUNIT AS 职务, r.RYBH AS 档案编号, r.WORKTIME AS 参加工作时间
            FROM RS_INFO r
            INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
            INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
            WHERE c.TID IN ({placeholders}) AND c.PID > 0
            ORDER BY r.XM
        """,
            unit_tids,
        )
        cols = [c[0] for c in cursor.description]
        persons = [dict(zip(cols, r)) for r in cursor.fetchall()]
        cursor.close()
    finally:
        if conn:
            conn.close()

    import openpyxl
    from datetime import datetime

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人员信息"
    headers = [
        "姓名",
        "性别",
        "出生年月",
        "政治面貌",
        "职务",
        "档案编号",
        "参加工作时间",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    for i, p in enumerate(persons, 2):
        for j, h in enumerate(headers, 1):
            ws.cell(row=i, column=j, value=p.get(h, ""))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"档案库存人员_{datetime.now().strftime('%Y%m%d')}.xlsx"
    resp = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return resp


@csrf_exempt
def detail_by_range_api(request):
    if request.method != "POST":
        return JsonResponse({"code": 405})
    try:
        data = json.loads(request.body)
        tids = data.get("tids", [])
        stat_type = data.get("type", "age")
        range_name = data.get("range", "")
    except:
        return JsonResponse({"code": 400})

    cache_key = f"archive_stats_detail_range:{stat_type}:{range_name}:{','.join(sorted(map(str, tids)))}"
    cached = cache.get(cache_key)
    if cached is not None:
        return JsonResponse({"code": 0, **cached})

    ranges = {
        "35岁以下": (0, 35),
        "36-40岁": (36, 40),
        "41-45岁": (41, 45),
        "46-50岁": (46, 50),
        "51-55岁": (51, 55),
        "56-60岁": (56, 60),
        "60岁以上": (61, 200),
        "5年以下": (0, 5),
        "6-10年": (6, 10),
        "11-15年": (11, 15),
        "16-20年": (16, 20),
        "21-25年": (21, 25),
        "26-30年": (26, 30),
        "30年以上": (31, 60),
        "未知": None,
    }
    rng = ranges.get(range_name)
    if rng is None and range_name != "未知":
        return JsonResponse({"code": 400, "msg": "无效范围"})

    conn = None
    try:
        conn = _get_conn()
        cursor = conn.cursor()
        rows = _get_unit_stats(cursor, tids)
        unit_tids = [str(r["tid"]) for r in rows]
        placeholders = ",".join(["?"] * len(unit_tids))

        if stat_type == "age":
            field = "r.CSNY"
        else:
            field = "r.WORKTIME"

        if range_name == "未知":
            where = f"AND ({field} IS NULL OR {field} = '' OR TRY_CAST(LEFT({field},4) AS INT) IS NULL)"
        else:
            lo, hi = rng
            where = f"AND TRY_CAST(LEFT({field},4) AS INT) BETWEEN YEAR(GETDATE())-{hi} AND YEAR(GETDATE())-{lo}"

        cursor.execute(
            f"SELECT COUNT(DISTINCT r.RSID) FROM RS_INFO r INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID WHERE c.TID IN ({placeholders}) AND c.PID > 0 {where}",
            unit_tids,
        )
        total = cursor.fetchone()[0]

        if total > 500:
            cursor.close()
            result = {"total": total, "data": [], "msg": "数据量超过500"}
            cache.set(cache_key, result, 3600)
            return JsonResponse({"code": 0, **result})

        cursor.execute(
            f"""
            SELECT DISTINCT r.XM, r.XB, r.CSNY, r.ZZMM, r.JOBUNIT, r.RYBH, r.WORKTIME
            FROM RS_INFO r
            INNER JOIN USERS_DEPARTMENT u ON r.RSID = u.RSID
            INNER JOIN BMGL c ON u.DEPARTMENTID = c.TID
            WHERE c.TID IN ({placeholders}) AND c.PID > 0 {where}
            ORDER BY r.XM
        """,
            unit_tids,
        )
        cols = [c[0] for c in cursor.description]
        persons = [dict(zip(cols, r)) for r in cursor.fetchall()]
        cursor.close()
        result = {"total": total, "data": persons}
        cache.set(cache_key, result, 3600)
        return JsonResponse({"code": 0, **result})
    except Exception as e:
        return JsonResponse({"code": 500, "msg": str(e)})
    finally:
        if conn:
            conn.close()
